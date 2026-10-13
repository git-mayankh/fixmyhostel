import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

DOCS_DIR = "/Users/mayankheda/Desktop/asd mpr/fixmyhostel/docs"
MD_PATH = os.path.join(DOCS_DIR, "MINI_PROJECT_REPORT.md")
DOCX_PATH = os.path.join(DOCS_DIR, "MINI_PROJECT_REPORT.docx")
HTML_PATH = os.path.join(DOCS_DIR, "MINI_PROJECT_REPORT.html")
SCREENSHOTS_DIR = os.path.join(DOCS_DIR, "screenshots")
DIAGRAMS_DIR = os.path.join(DOCS_DIR, "diagrams")

def set_cell_background(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_styled_table(doc, headers, rows):
    col_count = len(headers)
    table = doc.add_table(rows=len(rows) + 1, cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Format header
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title.strip()
        set_cell_background(hdr_cells[i], "1E3A8A")
        set_cell_margins(hdr_cells[i], 120, 120, 160, 160)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)

    # Format data rows
    for r_idx, row in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_col = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            if c_idx < col_count:
                row_cells[c_idx].text = val.strip()
                set_cell_background(row_cells[c_idx], bg_col)
                set_cell_margins(row_cells[c_idx], 80, 80, 120, 120)
                p = row_cells[c_idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph() # spacing

def add_code_block(doc, code_lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, 100, 100, 150, 150)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    code_text = "\n".join(code_lines)
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(15, 23, 42)
    doc.add_paragraph()

def add_formatted_text(p, text):
    # Regex split on bold (**text**) and code (`text`)
    pattern = r'(\*\*.*?\*\*|`.*?`|\*.*?\*)'
    tokens = re.split(pattern, text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            run = p.add_run(tok[2:-2])
            run.font.bold = True
        elif tok.startswith('`') and tok.endswith('`'):
            run = p.add_run(tok[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(180, 83, 9)
        elif tok.startswith('*') and tok.endswith('*'):
            run = p.add_run(tok[1:-1])
            run.font.italic = True
        else:
            p.add_run(tok)

def build_docx():
    print("Building DOCX report...")
    doc = docx.Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(4)

    with open(MD_PATH, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    i = 0
    in_code_block = False
    code_buffer = []
    code_lang = ""
    in_table = False
    table_headers = []
    table_rows = []

    mermaid_map = {
        0: os.path.join(DIAGRAMS_DIR, "fig_3_1_architecture.png"),
        1: os.path.join(DIAGRAMS_DIR, "fig_3_2_sequence.png"),
        2: os.path.join(DIAGRAMS_DIR, "fig_3_3_statemachine.png"),
        3: os.path.join(DIAGRAMS_DIR, "fig_3_4_er_diagram.png"),
        4: os.path.join(DIAGRAMS_DIR, "fig_5_1_devops_pipeline.png")
    }
    mermaid_idx = 0

    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        stripped = line.strip()

        # Handle code blocks
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                if code_lang == 'mermaid':
                    img_path = mermaid_map.get(mermaid_idx)
                    mermaid_idx += 1
                    if img_path and os.path.exists(img_path):
                        p_img = doc.add_paragraph()
                        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_img.paragraph_format.space_before = Pt(8)
                        p_img.paragraph_format.space_after = Pt(2)
                        run = p_img.add_run()
                        run.add_picture(img_path, width=Inches(6.0))
                else:
                    add_code_block(doc, code_buffer)
                code_buffer = []
                code_lang = ""
            else:
                in_code_block = True
                code_lang = stripped[3:].strip()
                code_buffer = []
            i += 1
            continue

        if in_code_block:
            code_buffer.append(line)
            i += 1
            continue

        # Handle tables
        if stripped.startswith('|') and stripped.endswith('|'):
            cells = [c.strip() for c in stripped.split('|')[1:-1]]
            # Check if separator row
            if all(re.match(r'^:?-+:?$', c) for c in cells):
                i += 1
                continue
            if not in_table:
                in_table = True
                table_headers = cells
                table_rows = []
            else:
                table_rows.append(cells)
            i += 1
            continue
        elif in_table:
            in_table = False
            add_styled_table(doc, table_headers, table_rows)
            table_headers = []
            table_rows = []

        # Empty line
        if not stripped:
            i += 1
            continue

        # Page dividers / breaks
        if stripped == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        # Image in markdown: ![caption](url)
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)$', stripped)
        if img_match:
            caption = img_match.group(1)
            raw_url = img_match.group(2)
            rel_path = raw_url.lstrip('./')
            full_img_path = os.path.join(DOCS_DIR, rel_path)
            if os.path.exists(full_img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(2)
                run = p_img.add_run()
                run.add_picture(full_img_path, width=Inches(6.0))
                
                # Caption
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(2)
                p_cap.paragraph_format.space_after = Pt(8)
                run_cap = p_cap.add_run(caption)
                run_cap.font.italic = True
                run_cap.font.size = Pt(9.5)
                run_cap.font.color.rgb = RGBColor(100, 116, 139)
            i += 1
            continue

        # Headings
        if stripped.startswith('# '):
            title_text = stripped[2:].strip()
            # If start of major chapter, add page break
            if title_text.startswith("CHAPTER") or title_text == "CERTIFICATE" or title_text == "ABSTRACT" or title_text == "TABLE OF CONTENTS":
                doc.add_page_break()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(title_text)
            run.font.name = "Calibri"
            run.font.size = Pt(17)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138) # Deep navy
            i += 1
            continue

        if stripped.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[3:].strip())
            run.font.name = "Calibri"
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42) # Slate dark
            i += 1
            continue

        if stripped.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[4:].strip())
            run.font.name = "Calibri"
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(51, 65, 85)
            i += 1
            continue

        # Bullet list
        if stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            add_formatted_text(p, stripped[2:].strip())
            i += 1
            continue

        # Numbered list
        num_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_match:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            add_formatted_text(p, num_match.group(2).strip())
            i += 1
            continue

        # Captions like *Figure 3.1: ...*
        if stripped.startswith('*Figure') and stripped.endswith('*'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(stripped[1:-1])
            run.font.italic = True
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(100, 116, 139)
            i += 1
            continue

        # Normal paragraph
        p = doc.add_paragraph()
        add_formatted_text(p, stripped)
        i += 1

    # End of while loop
    if in_table:
        add_styled_table(doc, table_headers, table_rows)

    doc.save(DOCX_PATH)
    print(f"DOCX created successfully at: {DOCX_PATH}")

def build_html():
    print("Building printable HTML report...")
    with open(MD_PATH, 'r', encoding='utf-8') as f:
        md_text = f.read()

    # Replace mermaid blocks with corresponding diagram images
    mermaid_replacements = [
        '<div class="figure-container"><img src="./diagrams/fig_3_1_architecture.png" alt="Figure 3.1: High-Level Architecture Diagram of FixMyHostel" /></div>',
        '<div class="figure-container"><img src="./diagrams/fig_3_2_sequence.png" alt="Figure 3.2: Complaint Lifecycle Sequence Diagram" /></div>',
        '<div class="figure-container"><img src="./diagrams/fig_3_3_statemachine.png" alt="Figure 3.3: Ticket Lifecycle State Machine" /></div>',
        '<div class="figure-container"><img src="./diagrams/fig_3_4_er_diagram.png" alt="Figure 3.4: Relational Database Schema & Data Models" /></div>',
        '<div class="figure-container"><img src="./diagrams/fig_5_1_devops_pipeline.png" alt="Figure 5.1: End-to-End DevOps Delivery Pipeline Flowchart" /></div>'
    ]

    parts = re.split(r'```mermaid[\s\S]*?```', md_text)
    rebuilt = parts[0]
    for idx, rep in enumerate(mermaid_replacements):
        if idx < len(parts) - 1:
            rebuilt += rep + parts[idx + 1]

    # Convert code blocks to <pre><code>
    def code_repl(m):
        lang = m.group(1) or ""
        code = m.group(2).replace('<', '&lt;').replace('>', '&gt;')
        return f'<pre><code class="lang-{lang}">{code}</code></pre>'
    rebuilt = re.sub(r'```(\w*)\n([\s\S]*?)```', code_repl, rebuilt)

    # Convert markdown images to HTML
    rebuilt = re.sub(r'!\[(.*?)\]\((.*?)\)', r'<div class="figure-container"><img src="\2" alt="\1" /><div class="figure-caption">\1</div></div>', rebuilt)

    # Convert markdown headers
    rebuilt = re.sub(r'^### (.*?)$', r'<h3>\1</h3>', rebuilt, flags=re.MULTILINE)
    rebuilt = re.sub(r'^## (.*?)$', r'<h2>\1</h2>', rebuilt, flags=re.MULTILINE)
    rebuilt = re.sub(r'^# (.*?)$', r'<h1>\1</h1>', rebuilt, flags=re.MULTILINE)

    # Convert markdown bold, italics, inline code
    rebuilt = re.sub(r'\*\*(.*?)\*\*', r'<strong>\1</strong>', rebuilt)
    rebuilt = re.sub(r'\*(.*?)\*', r'<em>\1</em>', rebuilt)
    rebuilt = re.sub(r'`([^`]+)`', r'<code>\1</code>', rebuilt)

    # Convert tables
    def table_repl(match):
        table_text = match.group(0).strip()
        lines = table_text.split('\n')
        html_tbl = ['<div class="table-container"><table>']
        hdr_done = False
        for line in lines:
            if not line.strip() or re.match(r'^\s*\|?\s*:?-+:?', line):
                continue
            cols = [c.strip() for c in line.strip().split('|')[1:-1]]
            if not hdr_done:
                html_tbl.append('<thead><tr>' + ''.join(f'<th>{c}</th>' for c in cols) + '</tr></thead><tbody>')
                hdr_done = True
            else:
                html_tbl.append('<tr>' + ''.join(f'<td>{c}</td>' for c in cols) + '</tr>')
        if hdr_done:
            html_tbl.append('</tbody></table></div>')
        return '\n'.join(html_tbl)

    rebuilt = re.sub(r'(\|[^\n]+\|\n)+', table_repl, rebuilt)

    # Convert lists and paragraphs
    lines = rebuilt.split('\n')
    final_lines = []
    in_ul = False
    in_ol = False

    for line in lines:
        s = line.strip()
        if s.startswith('- ') or s.startswith('* '):
            if not in_ul:
                final_lines.append('<ul>')
                in_ul = True
            final_lines.append(f'<li>{s[2:]}</li>')
            continue
        elif in_ul:
            final_lines.append('</ul>')
            in_ul = False

        num_m = re.match(r'^\d+\.\s+(.*)$', s)
        if num_m:
            if not in_ol:
                final_lines.append('<ol>')
                in_ol = True
            final_lines.append(f'<li>{num_m.group(1)}</li>')
            continue
        elif in_ol:
            final_lines.append('</ol>')
            in_ol = False

        if s == '---':
            final_lines.append('<hr class="page-break" />')
            continue

        if s and not s.startswith('<h') and not s.startswith('<table') and not s.startswith('<thead') and not s.startswith('<tbody') and not s.startswith('<tr') and not s.startswith('<div') and not s.startswith('</div') and not s.startswith('<pre') and not s.startswith('</pre') and not s.startswith('<code') and not s.startswith('</code') and not s.startswith('<li') and not s.startswith('<ul') and not s.startswith('</ul') and not s.startswith('<ol') and not s.startswith('</ol'):
            final_lines.append(f'<p>{s}</p>')
        else:
            final_lines.append(line)

    if in_ul:
        final_lines.append('</ul>')
    if in_ol:
        final_lines.append('</ol>')

    body_html = '\n'.join(final_lines)

    html_template = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>FixMyHostel - Mini Project Report</title>
<style>
  @page {{
    size: A4;
    margin: 20mm 18mm 20mm 18mm;
  }}
  body {{
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
    color: #1e293b;
    background: #ffffff;
    line-height: 1.6;
    margin: 0;
    padding: 30px;
    font-size: 11pt;
  }}
  .container {{
    max-width: 900px;
    margin: 0 auto;
  }}
  h1 {{
    color: #1e3a8a;
    font-size: 22pt;
    border-bottom: 2px solid #e2e8f0;
    padding-bottom: 8px;
    margin-top: 28px;
    page-break-after: avoid;
  }}
  h2 {{
    color: #0f172a;
    font-size: 15pt;
    margin-top: 22px;
    page-break-after: avoid;
  }}
  h3 {{
    color: #334155;
    font-size: 12.5pt;
    margin-top: 16px;
    page-break-after: avoid;
  }}
  p {{
    margin: 8px 0;
    text-align: justify;
  }}
  .figure-container {{
    text-align: center;
    margin: 20px 0 10px 0;
    page-break-inside: avoid;
  }}
  .figure-container img {{
    max-width: 95%;
    height: auto;
    border: 1px solid #cbd5e1;
    border-radius: 6px;
    box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.1);
  }}
  .figure-caption {{
    font-size: 9pt;
    font-style: italic;
    color: #64748b;
    margin-top: 6px;
  }}
  .table-container {{
    margin: 16px 0;
    overflow-x: auto;
    page-break-inside: avoid;
  }}
  table {{
    width: 100%;
    border-collapse: collapse;
    font-size: 9pt;
  }}
  th {{
    background-color: #1e3a8a;
    color: #ffffff;
    padding: 8px 12px;
    text-align: left;
    font-weight: 600;
  }}
  td {{
    padding: 8px 12px;
    border-bottom: 1px solid #e2e8f0;
  }}
  tr:nth-child(even) td {{
    background-color: #f8fafc;
  }}
  pre {{
    background-color: #f1f5f9;
    border: 1px solid #e2e8f0;
    border-radius: 6px;
    padding: 12px;
    font-family: ui-monospace, SFMono-Regular, Menlo, Monaco, Consolas, monospace;
    font-size: 8.5pt;
    overflow-x: auto;
    page-break-inside: avoid;
  }}
  code {{
    font-family: ui-monospace, SFMono-Regular, Menlo, Monaco, Consolas, monospace;
    background-color: #f1f5f9;
    padding: 2px 4px;
    border-radius: 4px;
    font-size: 9pt;
    color: #b45309;
  }}
  pre code {{
    background-color: transparent;
    padding: 0;
    color: #0f172a;
  }}
  ul, ol {{
    padding-left: 24px;
    margin: 8px 0;
  }}
  li {{
    margin-bottom: 4px;
  }}
  .page-break {{
    border: 0;
    border-top: 1px dashed #cbd5e1;
    margin: 30px 0;
    page-break-before: always;
  }}
  @media print {{
    body {{
      padding: 0;
    }}
    .page-break {{
      page-break-before: always;
      border: none;
      visibility: hidden;
    }}
    .no-print {{
      display: none;
    }}
  }}
  .print-bar {{
    position: sticky;
    top: 0;
    background: #1e293b;
    color: white;
    padding: 12px 24px;
    border-radius: 8px;
    display: flex;
    justify-content: space-between;
    align-items: center;
    margin-bottom: 24px;
    box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.1);
  }}
  .print-btn {{
    background: #3b82f6;
    color: white;
    border: none;
    padding: 8px 18px;
    border-radius: 6px;
    font-size: 10pt;
    font-weight: 600;
    cursor: pointer;
  }}
  .print-btn:hover {{
    background: #2563eb;
  }}
</style>
</head>
<body>
<div class="container">
  <div class="print-bar no-print">
    <div>
      <strong>FixMyHostel: Mini Project Academic Report</strong> — ASD&D Lab (Sem V-NEP)
    </div>
    <button class="print-btn" onclick="window.print()">Print / Save as PDF (Cmd + P)</button>
  </div>
  {body_html}
</div>
</body>
</html>
"""
    with open(HTML_PATH, 'w', encoding='utf-8') as f:
        f.write(html_template)
    print(f"HTML report created successfully at: {HTML_PATH}")

if __name__ == "__main__":
    build_docx()
    build_html()
    print("Report generation completed!")

