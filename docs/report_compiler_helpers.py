import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

DOCS_DIR = "/Users/mayankheda/Desktop/asd mpr/fixmyhostel/docs"
DOCX_PATH = os.path.join(DOCS_DIR, "MINI_PROJECT_REPORT.docx")
HTML_PATH = os.path.join(DOCS_DIR, "MINI_PROJECT_REPORT.html")
SCREENSHOTS_DIR = os.path.join(DOCS_DIR, "screenshots")
DIAGRAMS_DIR = os.path.join(DOCS_DIR, "diagrams")

def set_cell_background(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    cell._tc.get_or_add_tcPr().append(tcMar)

def set_cell_borders(cell, top="CCCCCC", bottom="CCCCCC", left="none", right="none"):
    tcBorders = parse_xml(f'''<w:tcBorders {nsdecls("w")}>
        <w:top w:val="{"single" if top != "none" else "none"}" w:sz="4" w:space="0" w:color="{top}"/>
        <w:left w:val="{"single" if left != "none" else "none"}" w:sz="4" w:space="0" w:color="{left}"/>
        <w:bottom w:val="{"single" if bottom != "none" else "none"}" w:sz="4" w:space="0" w:color="{bottom}"/>
        <w:right w:val="{"single" if right != "none" else "none"}" w:sz="4" w:space="0" w:color="{right}"/>
    </w:tcBorders>''')
    cell._tc.get_or_add_tcPr().append(tcBorders)

def add_lab_table(doc, headers, rows, col_widths=None):
    col_count = len(headers)
    table = doc.add_table(rows=len(rows) + 1, cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row (Exact Lab Manual style: light subtle gray #E2E8F0 or #F1F5F9, dark bold text)
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title.strip()
        set_cell_background(hdr_cells[i], "E2E8F0")
        set_cell_margins(hdr_cells[i], 120, 120, 150, 150)
        set_cell_borders(hdr_cells[i], top="94A3B8", bottom="94A3B8", left="CCCCCC", right="CCCCCC")
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)

    # Data Rows
    for r_idx, row in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_col = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            if c_idx < col_count:
                row_cells[c_idx].text = val.strip()
                set_cell_background(row_cells[c_idx], bg_col)
                set_cell_margins(row_cells[c_idx], 80, 80, 150, 150)
                set_cell_borders(row_cells[c_idx], top="E2E8F0", bottom="E2E8F0", left="E2E8F0", right="E2E8F0")
                p = row_cells[c_idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(30, 41, 59)

    # Width adjustment if provided
    if col_widths and len(col_widths) == col_count:
        for r in table.rows:
            for idx, width in enumerate(col_widths):
                r.cells[idx].width = width

    # Add spacing after table
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)

def add_code_terminal(doc, code_lines):
    # Terminal style: black/dark slate background matching Lab manual console outputs
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "0F172A")
    set_cell_margins(cell, 120, 120, 160, 160)
    set_cell_borders(cell, top="000000", bottom="000000", left="000000", right="000000")
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    
    code_text = "\n".join(code_lines)
    run = p.add_run(code_text)
    run.font.name = "Courier New"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(248, 250, 252)
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)

def add_body_p(doc, text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)

    # Token parsing for inline bold and code
    pattern = r'(\*\*.*?\*\*|`.*?`|\*.*?\*)'
    tokens = re.split(pattern, text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            run = p.add_run(tok[2:-2])
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            run.font.bold = True
        elif tok.startswith('`') and tok.endswith('`'):
            run = p.add_run(tok[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
        elif tok.startswith('*') and tok.endswith('*'):
            run = p.add_run(tok[1:-1])
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            run.font.italic = True
        else:
            run = p.add_run(tok)
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            if italic:
                run.font.italic = True
    return p

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_figure(doc, img_path, caption_text, width=Inches(5.8)):
    if not os.path.exists(img_path):
        print(f"Warning: image {img_path} not found!")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(img_path, width=width)
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(8)
    run_cap = p_cap.add_run(caption_text)
    run_cap.font.name = "Times New Roman"
    run_cap.font.size = Pt(10)
    run_cap.font.italic = True
    run_cap.font.color.rgb = RGBColor(51, 65, 85)

print("Helper definitions loaded successfully.")

